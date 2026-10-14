"""Downloads preserve user text, day boundaries and all scheduled rows."""
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from app.exporters import build_pdf, build_xlsx
from app.schemas import Activity, Problem, Resource


def _schedule():
    problem = Problem(
        scenario='life', name='我的学习与生活周计划',
        resources=[Resource(id='me', name='个人时间')],
        activities=[
            Activity(id='study', name='高等数学复习 <重点 & 难点>',
                     duration=3, resources=['me']),
            Activity(id='late', name='晚间阅读', duration=2, resources=['me'],
                     window_start=44, window_end=48),
            Activity(id='next', name='周二任务', duration=2, resources=['me']),
        ],
    )
    result = {'status': 'OPTIMAL', 'objective': {'value': 5},
              'metrics': {'hard_violations': 0, 'solve_time': 0.123},
              'assignments': [
                  {'activity_id': 'study', 'start': 18, 'end': 21},
                  {'activity_id': 'late', 'start': 46, 'end': 48},
                  {'activity_id': 'next', 'start': 48, 'end': 50},
              ]}
    return problem, result


def test_pdf_contains_chinese_user_text_and_literal_markup():
    problem, result = _schedule()
    payload = build_pdf(problem, result)
    reader = PdfReader(BytesIO(payload))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert problem.name in text
    assert '高等数学复习 <重点 & 难点>' in text
    assert '个人时间' in text
    assert 'OPTIMAL' in text and '0.123' in text
    assert '24:00' in text
    assert b'STSong-Light' in payload
    assert build_pdf(problem, result) == payload


def test_midnight_end_is_24_but_next_day_start_is_zero():
    problem, result = _schedule()
    sheet = load_workbook(BytesIO(build_xlsx(problem, result)))['调度明细']
    assert sheet.cell(3, 4).value == '周一'
    assert sheet.cell(3, 5).value == '23:00'
    assert sheet.cell(3, 6).value == '24:00'
    assert sheet.cell(4, 4).value == '周二'
    assert sheet.cell(4, 5).value == '00:00'


def test_excel_never_interprets_user_text_as_a_formula():
    problem, result = _schedule()
    problem.name = '=HYPERLINK("https://example.invalid", "方案")'
    problem.activities[0].name = '=1+1'
    problem.resources[0].name = '=SUM(1,2)'
    result['metrics']['=unsafe metric'] = '=1+2'
    workbook = load_workbook(BytesIO(build_xlsx(problem, result)), data_only=False)
    assert workbook['调度明细'].cell(2, 1).value == '=1+1'
    assert workbook['调度明细'].cell(2, 8).value == '=SUM(1,2)'
    assert workbook['求解概况'].cell(2, 2).value == problem.name
    assert all(cell.data_type != 'f' for sheet in workbook
               for row in sheet for cell in row)


def test_long_pdf_wraps_names_and_repeats_headers_without_dropping_tasks():
    resources = [Resource(id='me', name='个人时间')]
    problem = Problem(scenario='life', name='完整周计划', resources=resources,
                      activities=[Activity(id=f'a{i}', name=f'学习任务{i:03d}：' + '复杂内容' * 16,
                                           resources=['me']) for i in range(100)])
    result = {'status': 'FEASIBLE', 'assignments': [
        {'activity_id': f'a{i}', 'start': i * 2, 'end': i * 2 + 2} for i in range(100)]}
    reader = PdfReader(BytesIO(build_pdf(problem, result)))
    assert len(reader.pages) > 1
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert all(f'学习任务{i:03d}' in text for i in range(100))
    assert all('序号' in page.extract_text() for page in reader.pages)


def test_oversized_resource_cell_can_continue_on_another_pdf_page():
    problem, result = _schedule()
    problem.resources[0].name = '资源说明' * 1000
    reader = PdfReader(BytesIO(build_pdf(problem, result)))
    assert len(reader.pages) > 1
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert '周二任务' in text
