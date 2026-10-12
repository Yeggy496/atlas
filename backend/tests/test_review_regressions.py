"""Regression checks for the 2026-10-12 publication review."""
from fastapi.testclient import TestClient
from app.main import app
from app.schemas import Activity, Assignment, Constraint, Problem, Resource
from app.solver.explain import explain
from app.exporters import build_xlsx
from io import BytesIO
from openpyxl import load_workbook
import pytest
from app.adapters import ADAPTERS
from app.solver.diagnosis import add_demo_conflict
from app.solver.engine import solve


def test_diagnosis_requires_the_same_review_as_solve():
    with TestClient(app) as client:
        problem = client.get('/api/templates/lab').json()
        response = client.post('/api/diagnose', json={'problem': problem})
    assert response.status_code == 409


def test_repair_preserves_custom_weights():
    with TestClient(app) as client:
        problem = client.get('/api/templates/lab').json()
        constraint_id = problem['constraints'][0]['id']
        response = client.post('/api/repair', json={
            'problem': problem, 'reviewed': True,
            'constraint_id': constraint_id, 'weights': {'balance': 27},
        })
    assert response.status_code == 200
    result = response.json()['result']
    assert result['status'] in ('FEASIBLE', 'OPTIMAL')
    assert result['objective']['weights']['balance'] == 27


def test_explanation_does_not_apply_another_days_preference():
    problem = Problem(scenario='life', name='Day-specific preference',
        activities=[Activity(id='a', name='Study', days=[0], resources=['me'])],
        resources=[Resource(id='me', name='Me')],
        constraints=[Constraint(id='p', label='Tuesday morning', type='preferred_window',
            level='soft', target='a', day=1, start=16, end=20)])
    data = explain(problem, [Assignment(activity_id='a', start=16, end=18)], 'a')
    assert not any(reason['kind'] == 'preference' for reason in data['reasons'])


def test_excel_contains_selected_room_demands_and_awake_load():
    problem = Problem(scenario='life', name='Export resources',
        activities=[Activity(id='sleep', name='Sleep', category='sleep', duration=16,
                             window_start=0, window_end=16),
                    Activity(id='study', name='Study', duration=2,
                             room_options=['room'], demands={'device': 2})],
        resources=[Resource(id='room', name='Selected room'),
                   Resource(id='device', name='Equipment', capacity=2)])
    result = {'status': 'FEASIBLE', 'assignments': [
        {'activity_id': 'sleep', 'start': 0, 'end': 16},
        {'activity_id': 'study', 'start': 18, 'end': 20, 'room_id': 'room'}]}
    workbook = load_workbook(BytesIO(build_xlsx(problem, result)))
    resource_cell = workbook['调度明细'].cell(3, 8).value
    assert 'Selected room' in resource_cell and 'Equipment × 2' in resource_cell
    assert workbook['每日负载'].cell(2, 3).value == 1


@pytest.mark.parametrize('scenario', ['meeting', 'production'])
def test_new_scenarios_support_the_conflict_demo(scenario):
    problem = add_demo_conflict(ADAPTERS[scenario]())
    assert solve(problem, seconds=1)['status'] == 'INFEASIBLE'
