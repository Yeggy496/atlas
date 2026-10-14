"""Deterministic file import. Structured rows never pass through a model."""
import csv
from datetime import datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import math
import re
from uuid import uuid4

from openpyxl import load_workbook

from .schemas import Activity, Constraint, Problem


MAX_FILE_BYTES = 2_000_000
MAX_TEXT_CHARS = 200_000
PARSE_TEXT_CHARS = 4_000
TEXT_SUFFIXES = ('.txt',)
PDF_SUFFIXES = ('.pdf',)

HEADER_ALIASES = {
    '活动名称': 'name', '任务名称': 'name', '课程名称': 'name', '名称': 'name',
    '时长': 'duration_minutes', '时长(分钟)': 'duration_minutes',
    '时长（分钟）': 'duration_minutes', '持续时间(分钟)': 'duration_minutes',
    '持续时间（分钟）': 'duration_minutes', '分钟': 'duration_minutes',
    '星期': 'day', '周几': 'day', '日期': 'day',
    '开始时间': 'start', '固定时间': 'start',
    '资源': 'resources', '资源名称': 'resources', '资源id': 'resources',
    '类别': 'category', '分类': 'category',
}
SCENARIO_DEFAULTS = {
    'life': ('me', 'person', 'study'),
    'academic': ('class', 'class', 'course'),
    'lab': ('teamA', 'team', 'experiment'),
    'meeting': ('room_m', 'room', 'meeting'),
    'production': ('mill1', 'machine', 'production'),
}


def parse_clock(value):
    """Accept common Excel time cells without rounding unsupported precision."""
    hint = '开始时间请使用 00:00–23:30，按半小时填写（例如 19:00 或 19:30）。'
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond or value.minute not in (0, 30):
            raise ValueError(hint)
        return value.hour * 2 + value.minute // 30
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        # Excel can expose an unformatted clock as a fraction of a day.
        try:
            slots = Decimal(str(value)) * 48
        except InvalidOperation as exc:
            raise ValueError(hint) from exc
        nearest = slots.to_integral_value()
        if (not slots.is_finite() or not 0 <= nearest < 48
                or abs(slots - nearest) > Decimal('0.00000001')):
            raise ValueError(hint)
        return int(nearest)
    match = re.fullmatch(r'(\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d+))?)?', str(value).strip())
    if not match:
        raise ValueError(hint)
    hour, minute = int(match[1]), int(match[2])
    seconds, fraction = int(match[3] or 0), match[4] or ''
    if (not 0 <= hour <= 23 or minute not in (0, 30)
            or seconds or any(c != '0' for c in fraction)):
        raise ValueError(hint)
    return hour * 2 + minute // 30


def parse_integer(value, label):
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f'{label}须填写整数。') from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f'{label}须填写整数。')
    return int(number)


def _headers(values):
    values = list(values)
    while values and (values[-1] is None or not str(values[-1]).strip()):
        values.pop()
    result = []
    for value in values:
        key = str(value or '').strip().lower()
        result.append(HEADER_ALIASES.get(key, key))
    if len(set(result)) != len(result):
        raise ValueError('表头有重复列，请每种信息只保留一列。')
    if not {'name', 'duration_minutes'}.issubset(result):
        raise ValueError('文件缺少 name（活动名称）或 duration_minutes（时长/分钟）列，请下载模板。')
    return result


def _read_rows(filename, content):
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('上传文件最大为 2 MB。')
    if filename.lower().endswith('.csv'):
        raw = _decode_text(content)
        try:
            dialect = csv.Sniffer().sniff(raw[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(StringIO(raw), dialect=dialect)
        first = next(reader, None)
        if first is None:
            raise ValueError('文件为空，请下载模板并填写活动。')
        headers = _headers(first)
        rows = []
        for index, line in enumerate(reader, 2):
            if not any(str(value).strip() for value in line):
                continue
            if len(line) > len(headers):
                raise ValueError(f'第 {index} 行列数超过表头；名称中含逗号时请使用双引号包围。')
            rows.append((index, dict(zip(headers, line))))
        return rows
    if filename.lower().endswith('.xlsx'):
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            iterator = book.active.iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                raise ValueError('文件为空，请下载模板并填写活动。')
            headers = _headers(first)
            return [(index, dict(zip(headers, line))) for index, line in enumerate(iterator, 2)
                    if any(value is not None and str(value).strip() for value in line)]
        finally:
            book.close()
    raise ValueError('任务表支持 CSV 和 XLSX 模板。')


def _resource_ids(problem, cell):
    resources = {resource.id: resource for resource in problem.resources}
    names = {}
    for resource in problem.resources:
        names.setdefault(resource.name.strip(), []).append(resource.id)
    if cell is None or not str(cell).strip():
        default_id, kind, _ = SCENARIO_DEFAULTS[problem.scenario]
        if default_id in resources:
            return [default_id]
        candidates = [resource.id for resource in problem.resources if resource.type == kind]
        if len(candidates) == 1:
            return candidates
        if len(resources) == 1:
            return list(resources)
        raise ValueError('资源未填写且无法唯一确定。请填写当前资源的名称或 ID，多项用 | 分隔。')
    selected = []
    for token in re.split(r'[|；;、\n]+', str(cell)):
        token = token.strip()
        if not token:
            continue
        if token in resources:
            resource_id = token
        elif token in names and len(names[token]) == 1:
            resource_id = names[token][0]
        elif token in names:
            raise ValueError(f'资源「{token}」存在重名，请填写具体资源 ID。')
        else:
            raise ValueError(f'找不到资源「{token}」。请先在资源管理中添加，或使用现有资源名称/ID。')
        if resource_id not in selected:
            selected.append(resource_id)
    if not selected:
        raise ValueError('资源不能为空，请填写当前资源的名称或 ID。')
    return selected


def _day(value):
    chinese = re.fullmatch(r'(?:周|星期)([一二三四五六日天])', str(value).strip())
    if chinese:
        return '一二三四五六日'.index(chinese[1].replace('天', '日'))
    day = parse_integer(value, '日期')
    if not 1 <= day <= 7:
        raise ValueError('日期请填写 1–7（周一至周日），或「周三」这样的星期。')
    return day - 1


def _signature(activity, fixed):
    return (activity.name, activity.duration, activity.category,
            tuple(sorted(activity.resources)), tuple(sorted(activity.days)), fixed)


def import_table(problem: Problem, filename: str, content: bytes):
    rows = _read_rows(filename, content)
    if not rows or len(rows) > 100:
        raise ValueError('文件须包含 1–100 行活动。')
    if len(problem.activities) + len(rows) > 100:
        raise ValueError(f'当前有 {len(problem.activities)} 项活动，本次将新增 {len(rows)} 项；总数不能超过 100。请减少导入行数或先清空活动。')
    p = problem.model_copy(deep=True)
    fixed_by_activity = {c.target: (c.day, c.start) for c in p.constraints if c.type == 'fixed_time'}
    known = {_signature(activity, fixed_by_activity.get(activity.id)) for activity in p.activities}
    summary, warnings, duplicate_candidates = [], [], 0
    for index, row in rows:
        try:
            name = str(row.get('name') or '').strip()
            if not name or len(name) > 100:
                raise ValueError('活动名称须填写 1–100 个字符。')
            minutes = parse_integer(row.get('duration_minutes'), '时长')
            if minutes <= 0 or minutes > 1440 or minutes % 30:
                raise ValueError('时长须为 30–1440 分钟，且为 30 分钟整数倍。')
            ids = _resource_ids(p, row.get('resources'))
            activity = Activity(id='import_' + uuid4().hex[:10], name=name,
                duration=minutes // 30, resources=ids,
                category=str(row.get('category') or '').strip() or SCENARIO_DEFAULTS[p.scenario][2])
            day, start = row.get('day'), row.get('start')
            if day is not None and str(day).strip():
                activity.days = [_day(day)]
            fixed = None
            if start is not None and str(start).strip():
                if len(activity.days) != 1:
                    raise ValueError('固定时间必须同时指定 day（1–7）或星期。')
                slot = parse_clock(start)
                if slot + activity.duration > 48:
                    raise ValueError('固定活动的结束时间超过当天 24:00，请拆成两项活动。')
                fixed = (activity.days[0], slot)
                p.constraints.append(Constraint(id='fixed_' + activity.id,
                    label=f'{name} · 周{"一二三四五六日"[activity.days[0]]} {slot // 2:02}:{slot % 2 * 30:02} 开始',
                    type='fixed_time', target=activity.id, day=activity.days[0], start=slot, source='file'))
                activity.window_start, activity.window_end = 0, 48
            signature = _signature(activity, fixed)
            if signature in known:
                duplicate_candidates += 1
                warnings.append(f'第 {index} 行「{name}」与现有或本文件中的活动相似，已作为独立活动追加；请确认是否重复。')
            known.add(signature)
            p.activities.append(activity)
            summary.append(f'导入 {name} · {minutes} 分钟')
        except ValueError as exc:
            raise ValueError(f'第 {index} 行：{exc}') from exc
    if len(p.constraints) > 200:
        raise ValueError('导入后约束超过 200 条，请减少固定日程或先清理约束。')
    return {'problem': Problem.model_validate(p.model_dump()).model_dump(),
            'summary': summary, 'warnings': warnings, 'added_count': len(rows),
            'duplicate_candidates': duplicate_candidates, 'requires_review': True}


def extract_text(filename: str, content: bytes) -> dict:
    """Extract the entire supported document, leaving interpretation for review."""
    if len(content) > MAX_FILE_BYTES:
        raise ValueError('上传文件最大为 2 MB。')
    name = filename.lower()
    if name.endswith(TEXT_SUFFIXES):
        text, kind = _decode_text(content), 'txt'
    elif name.endswith(PDF_SUFFIXES):
        text, kind = _extract_pdf(content), 'pdf'
    else:
        raise ValueError('文本导入支持 TXT 和 PDF；结构化数据请使用 CSV / XLSX 模板。')
    text = text.strip()
    if not text:
        raise ValueError('未能从文件中提取到文字。扫描版 PDF 需要先做 OCR。')
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(f'提取的文字超过 {MAX_TEXT_CHARS:,} 字，请拆分文件后导入；没有截取或导入部分内容。')
    return {'kind': kind, 'text': text, 'chars': len(text), 'original_chars': len(text),
            'truncated': False, 'parse_limit': PARSE_TEXT_CHARS,
            'chunk_count': math.ceil(len(text) / PARSE_TEXT_CHARS)}


def _decode_text(content: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-8', 'gb18030'):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('无法识别文本编码，请另存为 UTF-8。')


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError('服务器未安装 PDF 解析组件（pypdf）。请改用 TXT，或让管理员安装依赖。') from exc
    try:
        reader = PdfReader(BytesIO(content))
        pages = [(page.extract_text() or '').strip() for page in reader.pages]
    except Exception as exc:
        raise ValueError('PDF 解析失败，请确认文件未加密且非扫描件。') from exc
    return '\n'.join(page for page in pages if page)
